"""Document Parser Tools — PDF, DOCX, Excel extraction.

Provides three tools:
- PDFReaderTool: extract text, metadata, page count from PDFs
- DocxReaderTool: extract text, headings, tables from DOCX files
- ExcelReaderTool: read sheets, headers, rows, stats from XLSX files
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any

from app.tools.base import BaseTool

logger = logging.getLogger(__name__)


class PDFReaderTool(BaseTool):
    """Extract text and metadata from PDF files."""

    def get_name(self) -> str:
        return "pdf_reader"

    def get_description(self) -> str:
        return (
            "Read and extract text from PDF files. Returns page text, "
            "metadata (title, author, page count), and optional full-text extraction."
        )

    def get_schema(self) -> dict:
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Path to PDF file"},
                "operation": {
                    "type": "string",
                    "enum": ["extract_text", "metadata", "page_count", "extract_page"],
                    "description": "Operation to perform",
                },
                "page": {"type": "integer", "description": "Page number for extract_page (1-indexed)"},
            },
            "required": ["file_path", "operation"],
        }

    async def execute(self, **kwargs: Any) -> dict:
        file_path = kwargs.get("file_path", "")
        operation = kwargs.get("operation", "extract_text")
        page_num = kwargs.get("page", 1)

        if not file_path:
            return {"error": "file_path is required"}
        if not Path(file_path).exists():
            return {"error": f"File not found: {file_path}"}

        try:
            import pymupdf
            doc = pymupdf.open(file_path)
        except ImportError:
            return {"error": "pymupdf not installed. Run: pip install pymupdf"}
        except Exception as e:
            return {"error": f"Failed to open PDF: {e}"}

        try:
            if operation == "metadata":
                metadata = doc.metadata or {}
                return {
                    "title": metadata.get("title", ""),
                    "author": metadata.get("author", ""),
                    "page_count": len(doc),
                    "file_size": Path(file_path).stat().st_size,
                }
            elif operation == "page_count":
                return {"page_count": len(doc)}
            elif operation == "extract_page":
                if page_num < 1 or page_num > len(doc):
                    return {"error": f"Page {page_num} out of range (1-{len(doc)})"}
                page = doc[page_num - 1]
                return {"page": page_num, "text": page.get_text()}
            else:
                full_text = ""
                for page in doc:
                    full_text += page.get_text()
                return {
                    "text": full_text[:50000],
                    "page_count": len(doc),
                    "char_count": len(full_text),
                    "truncated": len(full_text) > 50000,
                }
        finally:
            doc.close()


class DocxReaderTool(BaseTool):
    """Extract text and structure from DOCX files."""

    def get_name(self) -> str:
        return "docx_reader"

    def get_description(self) -> str:
        return (
            "Read and extract content from Word DOCX files. "
            "Returns text, headings, paragraphs, tables, and metadata."
        )

    def get_schema(self) -> dict:
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Path to DOCX file"},
                "operation": {
                    "type": "string",
                    "enum": ["extract_text", "headings", "tables", "metadata", "structure"],
                    "description": "Operation to perform",
                },
            },
            "required": ["file_path", "operation"],
        }

    async def execute(self, **kwargs: Any) -> dict:
        file_path = kwargs.get("file_path", "")
        operation = kwargs.get("operation", "extract_text")

        if not file_path:
            return {"error": "file_path is required"}
        if not Path(file_path).exists():
            return {"error": f"File not found: {file_path}"}

        try:
            from docx import Document
            doc = Document(file_path)
        except ImportError:
            return {"error": "python-docx not installed. Run: pip install python-docx"}
        except Exception as e:
            return {"error": f"Failed to open DOCX: {e}"}

        if operation == "metadata":
            props = doc.core_properties
            return {
                "title": props.title or "",
                "author": props.author or "",
                "created": str(props.created) if props.created else "",
                "modified": str(props.modified) if props.modified else "",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
            }
        elif operation == "headings":
            headings = [
                {"text": p.text, "level": p.style.name}
                for p in doc.paragraphs
                if p.style.name.startswith("Heading")
            ]
            return {"headings": headings}
        elif operation == "tables":
            tables = []
            for i, table in enumerate(doc.tables):
                rows = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                tables.append({"index": i, "rows": rows})
            return {"tables": tables}
        elif operation == "structure":
            elements = []
            for p in doc.paragraphs:
                if p.text.strip():
                    elements.append({
                        "type": "heading" if p.style.name.startswith("Heading") else "paragraph",
                        "style": p.style.name,
                        "text": p.text[:500],
                    })
            return {"elements": elements[:200], "total": len(elements)}
        else:
            text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n".join(text_parts)
            return {
                "text": full_text[:50000],
                "paragraph_count": len(doc.paragraphs),
                "char_count": len(full_text),
                "truncated": len(full_text) > 50000,
            }


class ExcelReaderTool(BaseTool):
    """Read and analyze Excel/XLSX/CSV files."""

    def get_name(self) -> str:
        return "excel_reader"

    def get_description(self) -> str:
        return (
            "Read and analyze spreadsheet files (XLSX, CSV). "
            "Returns sheet names, headers, data rows, and basic statistics."
        )

    def get_schema(self) -> dict:
        return {
            "type": "object",
            "properties": {
                "file_path": {"type": "string", "description": "Path to spreadsheet file"},
                "operation": {
                    "type": "string",
                    "enum": ["sheets", "headers", "read", "stats", "search"],
                    "description": "Operation to perform",
                },
                "sheet": {"type": "string", "description": "Sheet name (XLSX only, defaults to first)"},
                "max_rows": {"type": "integer", "description": "Max rows to read (default 100)"},
                "query": {"type": "string", "description": "Search query for search operation"},
            },
            "required": ["file_path", "operation"],
        }

    async def execute(self, **kwargs: Any) -> dict:
        file_path = kwargs.get("file_path", "")
        operation = kwargs.get("operation", "read")
        sheet_name = kwargs.get("sheet")
        max_rows = kwargs.get("max_rows", 100)
        query = kwargs.get("query", "")

        if not file_path:
            return {"error": "file_path is required"}
        if not Path(file_path).exists():
            return {"error": f"File not found: {file_path}"}

        ext = Path(file_path).suffix.lower()

        if ext == ".csv":
            return self._read_csv(file_path, operation, max_rows, query)
        elif ext in (".xlsx", ".xls"):
            return self._read_xlsx(file_path, operation, sheet_name, max_rows, query)
        else:
            return {"error": f"Unsupported format: {ext}. Use CSV or XLSX."}

    def _read_csv(self, path: str, operation: str, max_rows: int, query: str) -> dict:
        try:
            with open(path, encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                headers = next(reader, [])
                rows = []
                for i, row in enumerate(reader):
                    if i >= max_rows:
                        break
                    if query and not any(query.lower() in str(c).lower() for c in row):
                        continue
                    rows.append(row)
                return {
                    "headers": headers,
                    "rows": rows,
                    "row_count": len(rows),
                    "file": path,
                }
        except Exception as e:
            return {"error": f"CSV read error: {e}"}

    def _read_xlsx(self, path: str, operation: str, sheet_name: str | None, max_rows: int, query: str) -> dict:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except ImportError:
            return {"error": "openpyxl not installed. Run: pip install openpyxl"}
        except Exception as e:
            return {"error": f"Failed to open XLSX: {e}"}

        try:
            if operation == "sheets":
                return {"sheets": wb.sheetnames}

            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            headers = [str(c) if c is not None else "" for c in (all_rows[0] if all_rows else [])]

            if operation == "headers":
                return {"headers": headers, "sheet": ws.title}

            data_rows = []
            for row in all_rows[1:]:
                vals = [str(c) if c is not None else "" for c in row]
                if query and not any(query.lower() in v.lower() for v in vals):
                    continue
                data_rows.append(vals)
                if len(data_rows) >= max_rows:
                    break

            if operation == "stats":
                numeric_cols = []
                for col_idx, header in enumerate(headers):
                    nums = []
                    for row in data_rows:
                        try:
                            nums.append(float(row[col_idx]))
                        except (ValueError, IndexError):
                            pass
                    if nums:
                        numeric_cols.append({
                            "column": header,
                            "count": len(nums),
                            "min": min(nums),
                            "max": max(nums),
                            "avg": sum(nums) / len(nums),
                        })
                return {"headers": headers, "row_count": len(data_rows), "numeric_stats": numeric_cols}

            if operation == "search":
                return {"headers": headers, "rows": data_rows, "row_count": len(data_rows), "query": query}

            return {
                "headers": headers,
                "rows": data_rows,
                "row_count": len(data_rows),
                "sheet": ws.title,
                "truncated": len(all_rows) - 1 > max_rows,
            }
        finally:
            wb.close()
